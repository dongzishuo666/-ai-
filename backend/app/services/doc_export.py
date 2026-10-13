# ============================================================
# 文件：backend/app/services/doc_export.py
# 说明：把功能一产出的 Markdown 成品转换为可下载的 Word 文档。
#   支持块级：标题 / 段落 / 无序列表 / 有序列表 / 引用 / 代码块 / 表格 / 分隔线
#   支持行内：粗体 / 斜体 / 行内代码
#   中文字体走 eastAsia 设置，避免 Word 里中文回退成宋体以外的乱字体。
# ============================================================

from __future__ import annotations

import io
import re
from typing import List, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# 中文字体（Windows 与 WPS 通用）；等宽字体用于代码
_CJK_FONT = "微软雅黑"
_MONO_FONT = "Consolas"

_CODE_BG = "F5F7FA"          # 代码块底纹
_QUOTE_COLOR = RGBColor(0x60, 0x60, 0x60)
_BODY_SIZE = Pt(11)

# 块级正则
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*+]\s+(.*)$")
_ORDERED_RE = re.compile(r"^(\d+)[.)]\s+(.*)$")
_QUOTE_RE = re.compile(r"^>\s?(.*)$")
_TABLE_ROW_RE = re.compile(r"^\|.*\|\s*$")
_HR_RE = re.compile(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$")

# 行内正则：粗体 / 行内代码 / 斜体
_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\n]+\*)")


# ============================================================
# 底层工具
# ============================================================

def _apply_font(run, font_name: str = _CJK_FONT) -> None:
    """同时设置 ascii/hAnsi/eastAsia，保证中英文都落在同一个字体上。"""
    run.font.name = font_name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), font_name)


def _shade(paragraph, hex_fill: str) -> None:
    """给段落加底纹（代码块用）。"""
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def _bottom_border(paragraph) -> None:
    """给段落加下边框（分隔线用）。"""
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D0D7DE")
    borders.append(bottom)
    pPr.append(borders)


def _add_inline(paragraph, text: str) -> None:
    """解析行内标记，按段写入 run。"""
    for seg in _INLINE_RE.split(text):
        if not seg:
            continue
        if len(seg) > 4 and seg.startswith("**") and seg.endswith("**"):
            run = paragraph.add_run(seg[2:-2])
            run.bold = True
            _apply_font(run)
        elif len(seg) > 2 and seg.startswith("`") and seg.endswith("`"):
            run = paragraph.add_run(seg[1:-1])
            run.font.size = Pt(10)
            _apply_font(run, _MONO_FONT)
        elif len(seg) > 2 and seg.startswith("*") and seg.endswith("*"):
            run = paragraph.add_run(seg[1:-1])
            run.italic = True
            _apply_font(run)
        else:
            run = paragraph.add_run(seg)
            _apply_font(run)


def _new_body(doc: Document, text: str = "", indent_cm: float = 0.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.4
    if indent_cm:
        pf.left_indent = Cm(indent_cm)
    if text:
        _add_inline(p, text)
    return p


# ============================================================
# 块级渲染
# ============================================================

def _add_code_block(doc: Document, lines: List[str]) -> None:
    body = "\n".join(lines).rstrip("\n")
    if not body.strip():
        return
    # 多行代码拆成多段，靠底纹连成一块
    for line in body.split("\n"):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.left_indent = Cm(0.5)
        pf.line_spacing = 1.15
        run = p.add_run(line if line else " ")
        run.font.size = Pt(9.5)
        _apply_font(run, _MONO_FONT)
        _shade(p, _CODE_BG)
    # 代码段后留一点间距
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)


def _split_table_row(line: str) -> List[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _is_table_separator(line: str) -> bool:
    s = line.strip().replace("|", "")
    return bool(s) and set(s) <= set("-: ")


def _add_table(doc: Document, rows: List[List[str]]) -> None:
    if not rows:
        return
    col_count = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            _add_inline(p, row[c_idx] if c_idx < len(row) else "")
            if r_idx == 0:  # 表头加粗
                for run in p.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


# ============================================================
# 主入口
# ============================================================

def markdown_to_docx(md_text: str, title: str = "") -> io.BytesIO:
    """
    把 Markdown 文本转成 .docx。

    参数：
        md_text  AI 产出的成品正文（Markdown）
        title    文档标题（一般为原文件名），留空则不生成标题

    返回：
        指向 docx 二进制内容的 BytesIO（游标已在 0）
    """
    doc = Document()

    # 正文默认样式：中文落到微软雅黑
    normal = doc.styles["Normal"]
    normal.font.name = _CJK_FONT
    normal.font.size = _BODY_SIZE
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), _CJK_FONT
    )

    if title:
        heading = doc.add_heading(title, level=0)
        for run in heading.runs:
            _apply_font(run)

    lines = (md_text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0
    total = len(lines)

    while i < total:
        raw = lines[i]
        line = raw.strip()

        # 围栏代码块
        if line.startswith("```"):
            i += 1
            buf: List[str] = []
            while i < total and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # 跳过收尾围栏
            _add_code_block(doc, buf)
            continue

        if not line:
            i += 1
            continue

        # 分隔线
        if _HR_RE.match(line):
            hr = doc.add_paragraph()
            hr.paragraph_format.space_after = Pt(6)
            _bottom_border(hr)
            i += 1
            continue

        # 标题
        m = _HEADING_RE.match(line)
        if m:
            level = min(len(m.group(1)), 6)
            heading = doc.add_heading(level=level)
            _add_inline(heading, m.group(2).strip())
            for run in heading.runs:
                _apply_font(run)
            i += 1
            continue

        # 表格（连续 | 行）
        if _TABLE_ROW_RE.match(line) and not _is_table_separator(line):
            rows: List[List[str]] = []
            while i < total and _TABLE_ROW_RE.match(lines[i].strip()):
                cur = lines[i].strip()
                if not _is_table_separator(cur):
                    rows.append(_split_table_row(cur))
                i += 1
            _add_table(doc, rows)
            continue

        # 引用
        m = _QUOTE_RE.match(line)
        if m:
            p = _new_body(doc, "", indent_cm=0.75)
            if p.runs:
                for run in p.runs:
                    run.font.color.rgb = _QUOTE_COLOR
            _add_inline(p, m.group(1).strip())
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = _QUOTE_COLOR
            i += 1
            continue

        # 列表：手动加前缀而不用 List Bullet 样式，避免模板缺样式报错
        m = _BULLET_RE.match(line)
        if m:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.75)
            pf.space_after = Pt(3)
            run = p.add_run("• ")
            _apply_font(run)
            _add_inline(p, m.group(1).strip())
            i += 1
            continue

        m = _ORDERED_RE.match(line)
        if m:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.75)
            pf.space_after = Pt(3)
            run = p.add_run(f"{m.group(1)}. ")
            _apply_font(run)
            _add_inline(p, m.group(2).strip())
            i += 1
            continue

        # 普通段落：把后续非空行并入同一段（Markdown 软换行）
        buf = [line]
        i += 1
        while i < total:
            nxt = lines[i].strip()
            if (not nxt or _HEADING_RE.match(nxt) or _BULLET_RE.match(nxt)
                    or _ORDERED_RE.match(nxt) or _QUOTE_RE.match(nxt)
                    or _TABLE_ROW_RE.match(nxt)
                    or nxt.startswith("```") or _HR_RE.match(nxt)):
                break
            buf.append(nxt)
            i += 1
        _new_body(doc, " ".join(buf))
        continue

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
