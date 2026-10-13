# ============================================================
# 文件：backend/app/services/doc_parser.py
# 说明：文档解析。4 种格式 → 纯文本。
# 约束：全部同步 IO 操作，调用方用 asyncio.to_thread 包装。
# ============================================================

from __future__ import annotations

import io
import logging
from typing import Callable, Dict

import docx                # python-docx
import PyPDF2

logger = logging.getLogger("app.doc_parser")


# 支持的扩展名（与 DDL CHECK 一致）
SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt", "md"}


class DocParseError(Exception):
    """文档解析失败。"""
    pass


# ============================================================
# 各格式解析器
# ============================================================

def _parse_pdf(content: bytes) -> str:
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise DocParseError("PDF 已加密，无法解析")
        parts = []
        for page in reader.pages:
            try:
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(text)
            except Exception as e:  # noqa: BLE001
                logger.warning("PDF 单页解析失败，跳过：%s", e)
        return "\n".join(parts)
    except DocParseError:
        raise
    except Exception as e:
        raise DocParseError(f"PDF 解析异常: {e}") from e


def _parse_docx(content: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        # 补充表格文本
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    table_texts.append(" | ".join(cells))
        parts = paragraphs + table_texts
        return "\n".join(parts)
    except Exception as e:
        raise DocParseError(f"DOCX 解析异常: {e}") from e


def _parse_txt(content: bytes) -> str:
    return _decode_text(content)


def _parse_md(content: bytes) -> str:
    return _decode_text(content)


def _decode_text(content: bytes) -> str:
    """文本解码：优先 utf-8，失败退 gbk，再失败忽略错误。"""
    if not content:
        return ""
    for encoding in ("utf-8", "gbk", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


# ============================================================
# 分派表
# ============================================================

_PARSERS: Dict[str, Callable[[bytes], str]] = {
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "txt": _parse_txt,
    "md": _parse_md,
}


# ============================================================
# 对外入口
# ============================================================

def detect_doc_type(filename: str) -> str:
    """
    从文件名后缀识别文档类型。
    返回：'pdf' / 'docx' / 'txt' / 'md'
    异常：DocParseError（不支持格式）
    """
    if not filename:
        raise DocParseError("文件名为空")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocParseError(
            f"不支持的格式 .{ext}，仅支持 PDF/DOCX/TXT/MD"
        )
    return ext


def parse_document(content: bytes, doc_type: str) -> str:
    """
    解析文档为纯文本。
    异常：DocParseError
    """
    parser = _PARSERS.get(doc_type)
    if parser is None:
        raise DocParseError(f"无解析器处理类型: {doc_type}")
    text = parser(content)
    # 归一化：压缩连续空行
    lines = [ln.rstrip() for ln in text.splitlines()]
    normalized = "\n".join(lines)
    while "\n\n\n" in normalized:
        normalized = normalized.replace("\n\n\n", "\n\n")
    return normalized.strip()
